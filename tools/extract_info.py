"""Extract the 'Divinity: Original Sin 2 - Info (Shared)' workbook into JSON tables.
Usage: python3 -I extract_info.py <xlsx> <out.json>
Each table: list of row dicts. Cell hover-notes are kept under row['_notes'][field].
"""
import sys, json, re
import openpyxl
from openpyxl.utils import column_index_from_string as CI

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

def clean(v):
    if v is None: return None
    if isinstance(v, float):
        return int(v) if v.is_integer() else round(v, 4)
    if isinstance(v, str):
        v = v.strip()
        if v == '': return None
        try:
            f = float(v)
            if re.fullmatch(r'-?\d+(\.\d+)?', v):
                return int(f) if f.is_integer() else f
        except ValueError:
            pass
    return v

def note(c):
    if not c.comment: return None
    t = c.comment.text.strip()
    # Google exports prefix with author lines like "Author:\n" / "-Name"
    if t.startswith('[Threaded comment]'):
        t = re.sub(r'^\[Threaded comment\].*?\nComment:\n', '', t, flags=re.S)
        t = re.sub(r'\n(Reply|Comment):\n', '\n— ', t)
        t = re.sub(r'\n— \t', '\n— ', t)
    t = re.sub(r'\n\s*-[A-Z][\w .]*$', '', t).strip()
    return t or None

def table(name, cols, first=2, last=None, key='A', stop_blank=False):
    """cols: {field: column letter}. Rows where the key column is empty are skipped."""
    ws = wb[name]
    out = []
    last = last or ws.max_row
    for r in range(first, last + 1):
        kv = ws[f'{key}{r}'].value
        if kv in (None, ''):
            if stop_blank: break
            continue
        row, notes = {}, {}
        for f, col in cols.items():
            c = ws[f'{col}{r}']
            v = clean(c.value)
            if v is not None: row[f] = v
            n = note(c)
            if n: notes[f] = n
        if notes: row['_notes'] = notes
        out.append(row)
    return out

T = {}

# ---------- Recipes ----------
ws = wb['Recipes']
rec = []
for r in range(2, ws.max_row + 1):
    res = clean(ws[f'A{r}'].value)
    if not res: continue
    g = lambda col: clean(ws[f'{col}{r}'].value)
    row = {'result': res}
    for f, col in [('amount','B'),('effect','E'),('result2','F'),('amount2','G'),('category','H'),
                   ('subcat','I'),('station','J'),('autolevel','Z')]:
        v = g(col)
        if v is not None and v != 'V': row[f] = v
    ings = []
    for vcol, icol, tcol in [('K','L','M'),('N','O','P'),('Q','R','S'),('T','U','V'),('W','X','Y')]:
        name = g(icol)
        if name:
            ings.append({'name': name, 'value': g(vcol), 'mode': g(tcol)})
    row['ingredients'] = ings
    notes = {}
    for col, f in [('A','result'),('E','effect'),('F','result2'),('L','ing1'),('O','ing2'),('R','ing3'),('U','ing4'),('X','ing5')]:
        n = note(ws[f'{col}{r}'])
        if n: notes[f] = n
    if notes: row['_notes'] = notes
    rec.append(row)
T['recipes'] = rec

T['ingredients'] = table('Ingredients', {'name':'A','value':'B','category':'C','notes':'D'})

# ---------- Skills (this sheet's version; DE flag in col M) ----------
T['skills'] = table('Skills', {'name':'A','ability1':'B','ability2':'C','level':'D','ap':'E','source':'F',
    'cooldown':'G','range':'H','radius':'I','resist':'J','description':'K','numbers':'L','de':'M',
    'weaponReq':'N','scaling':'O','statusRef':'P','status1':'Q','status2':'R','status3':'S','numType':'T'})

# ---------- 2-ability (combo) skill matrix ----------
ws = wb['2-Ability Skills']
elems = [ws.cell(1, c).value for c in range(2, 6)]
combos = []
for r in range(2, 8):
    other = ws.cell(r, 1).value
    for i, el in enumerate(elems):
        sk = clean(ws.cell(r, 2 + i).value)
        if sk:
            row = {'skill': sk, 'element': el, 'other': other}
            n = note(ws.cell(r, 2 + i))
            if n: row['_notes'] = {'skill': n}
            combos.append(row)
T['combos'] = combos

T['surfaces'] = table('SurfacesClouds', {'name':'A','kind':'B','description':'C','duration':'D','resist':'E',
    'interactions':'F','vaporizeSpike':'G'})
T['statuses'] = table('Status', {'name':'A','resist':'B','effect':'C'})

# ---------- Runes ----------
ws = wb['Runes']
runes = {'boosts': [], 'mythic': [], 'power': [], 'sizes': [], 'sourceOrb': None}
section = 'boosts'
for r in range(2, ws.max_row + 1):
    a = clean(ws[f'A{r}'].value)
    if a is None:
        # size table lives in F:I independent of A
        pass
    if isinstance(a, str) and a.startswith('v '):
        section = 'mythic' if 'Mythic' in a else 'power'; continue
    if a == 'Source Orb':
        runes['sourceOrb'] = {'weapon': clean(ws[f'B{r}'].value), 'armor': clean(ws[f'C{r}'].value), 'amulet': clean(ws[f'D{r}'].value)}
        continue
    if a:
        row = {'rune': a, 'weapon': clean(ws[f'B{r}'].value)}
        if section != 'power':
            row['armor'] = clean(ws[f'C{r}'].value); row['amulet'] = clean(ws[f'D{r}'].value)
        notes = {k: note(ws[f'{c}{r}']) for k, c in [('rune','A'),('weapon','B'),('armor','C'),('amulet','D')] if note(ws[f'{c}{r}'])}
        if notes: row['_notes'] = notes
        runes[section].append(row)
    s = clean(ws[f'F{r}'].value)
    if s:
        runes['sizes'].append({'size': s, 'weapon': clean(ws[f'G{r}'].value), 'armor': clean(ws[f'H{r}'].value), 'amulet': clean(ws[f'I{r}'].value)})
T['runes'] = runes

# ---------- Summons (base + infused columns) ----------
T['summons'] = table('Summons', {'name':'A','level':'B','summoningLvl':'C','vitality':'D','physArmor':'E','magArmor':'F',
    'initiative':'G','damage':'H','dmgType':'I','airRes':'J','earthRes':'K','fireRes':'L','poisonRes':'M','waterRes':'N',
    'abilities':'O','skills':'P','immunities':'Q','preDE':'R','infusion':'S','infType':'T','infAirRes':'U','infEarthRes':'V',
    'infFireRes':'W','infPoisonRes':'X','infWaterRes':'Y','infAbilities':'Z','infSkills':'AA','infImmunities':'AB','infSpecial':'AC'})

# ---------- V: crafted item gold values by level ----------
ws = wb['V']
vals = []
for r in range(2, ws.max_row + 1):
    a = clean(ws[f'A{r}'].value)
    if not a: continue
    if isinstance(a, str) and a.startswith('v '): break  # testing block follows
    by = {}
    for c in range(2, 22):
        v = clean(ws.cell(r, c).value)
        if v is not None: by[str(c - 1)] = v
    vals.append({'item': a, 'byLevel': by})
T['values'] = vals

# ---------- Attributes (+ "Other Stats" section) ----------
ws = wb['Attributes']
attrs, other = [], []; sect = attrs
for r in range(2, ws.max_row + 1):
    a = clean(ws[f'A{r}'].value)
    if not a: continue
    if a == 'Other Stats': sect = other; continue
    row = {'name': a, 'effect': clean(ws[f'B{r}'].value), 'notes': clean(ws[f'C{r}'].value)}
    row = {k: v for k, v in row.items() if v is not None}
    nn = {k: note(ws[f'{c}{r}']) for k, c in [('name','A'),('effect','B'),('notes','C')] if note(ws[f'{c}{r}'])}
    if nn: row['_notes'] = nn
    sect.append(row)
T['attributes'] = attrs
T['stats'] = other

T['abilities'] = table('Abilities', {'name':'A','type':'B','effect':'C','notes':'D'})
T['talents'] = table('Talents', {'name':'A','description':'B','requirements':'C'})
T['talentPoints'] = table('Talents', {'level':'E','total':'F'}, key='E')
T['raceTraits'] = table('Race Traits', {'race':'A','trait':'B','description':'C'})
T['classes'] = table('Classes', {'name':'A','ability1':'B','ability2':'C','civil':'D','talent':'E','str':'F','fin':'G',
    'int':'H','con':'I','wit':'J','skill1':'K','skill2':'L','skill3':'M','weapons':'N','shield':'O','extraItems':'P','extraValue':'Q'})

# ---------- Weapons: base types + comparison rows ----------
w = table('Weapons', {'name':'A','hands':'B','type':'C','attribute':'D','range':'E','d1':'F','d5':'G','d10':'H',
    'd15':'I','d20':'J','d22':'K','notes':'M'})
T['weapons'] = [x for x in w if isinstance(x.get('hands'), int)]
T['weaponCompare'] = [x for x in w if 'hands' not in x]

# ---------- Armor ----------
arm = table('Armor', {'material':'A','slot':'B','attribute':'C','p1':'E','m1':'F','p10':'H','m10':'I','p20':'K','m20':'L','notes':'M'})
T['armor'] = [x for x in arm if x.get('slot') not in ('No shield', 'w/ Shield') and 'p1' in x]
ws = wb['Armor']; sets = []
for r in range(2, ws.max_row + 1):
    b = clean(ws[f'B{r}'].value)
    if b in ('No shield', 'w/ Shield'):
        v = [clean(ws[f'{c}{r}'].value) for c in ('E','F','H','I','K','L')]
        sets.append({'set': clean(ws[f'A{r}'].value), 'shield': b == 'w/ Shield', 'p1': v[0], 'm1': v[1], 'p10': v[2], 'm10': v[3], 'p20': v[4], 'm20': v[5]})
T['armorSets'] = sets

# ---------- Gear Stats: possible rolls per slot ----------
ws = wb['Gear Stats']
slots = [ws.cell(1, c).value for c in range(1, 15)]
rolls = []
rx = re.compile(r'^(?P<stat>.+?)[:]?\s+(?P<min>-?\d+(?:\.\d+)?)(?:-(?P<max>\d+(?:\.\d+)?))?(?P<pct>%)?$')
for c, slot in enumerate(slots, start=1):
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, c); v = clean(cell.value)
        if v is None: continue
        v = str(v)
        row = {'slot': slot, 'text': v}
        if v.startswith('Skill: '):
            row.update(kind='skill', stat=v[7:])
        elif v.startswith('Inflict: '):
            row.update(kind='inflict', stat=v[9:])
        elif v == 'Rune Slot':
            row.update(kind='rune', stat='Rune Slot')
        else:
            m = rx.match(v)
            if m:
                stat = m['stat'].replace('Critical Chance', 'Crit Chance')
                stat = {'Wit': 'Wits'}.get(stat, stat)
                kind = 'set' if stat.startswith('Set ') else 'stat'
                row.update(kind=kind, stat=stat, min=float(m['min']), max=float(m['max'] or m['min']), pct=bool(m['pct']))
                for k in ('min', 'max'):
                    if row[k].is_integer(): row[k] = int(row[k])
            else:
                row.update(kind='other', stat=v)
        n = note(cell)
        if n: row['_notes'] = {'text': n}
        rolls.append(row)
T['gearRolls'] = rolls

# ---------- Rarity multipliers ----------
ws = wb['Rarity Multipliers']; rar = []
for r in range(2, ws.max_row + 1):
    a = clean(ws[f'A{r}'].value)
    if not a: continue
    g = lambda col: clean(ws[f'{col}{r}'].value)
    row = {'type': a, 'white': g('B'), 'uncommon': g('C'), 'rare': g('E'), 'epic': g('G'), 'legendary': g('I'), 'divine': g('K'),
           'xUncommon': g('D'), 'xRare': g('F'), 'xEpic': g('H'), 'xLegendary': g('J'), 'xDivine': g('L')}
    row = {k: v for k, v in row.items() if v is not None}
    nn = {}
    for col in 'ABCDEFGHIJKL':
        n = note(ws[f'{col}{r}'])
        if n: nn[col] = n
    if nn: row['_notes'] = nn
    rar.append(row)
T['rarity'] = rar
hdr_notes = {c.column_letter: note(c) for c in ws[1] if note(c)}
T['_sheetNotes'] = {'rarityHeader': hdr_notes}

T['bugs'] = table('Bugs', {'bug':'A','conditions':'B','steps':'C','notes':'D'})

# index tab (tab descriptions + DE status)
T['tabs'] = table('|', {'tab':'A','description':'B','status':'C'}, first=8)

meta = {'source': 'Divinity: Original Sin 2 - Info (Shared) (Google Sheet 1CyOtdPIicJMgAK8RtAaFMNpBqoZDDzFp-IKY0JAWr38)',
        'counts': {k: (len(v) if isinstance(v, list) else None) for k, v in T.items()}}
T['_meta'] = meta
json.dump(T, open(sys.argv[2], 'w'), ensure_ascii=False, indent=0)
print(json.dumps(meta['counts']))
