import openpyxl, json, hashlib, io, base64, sys
from PIL import Image
SRC=sys.argv[1]; OUT=sys.argv[2]
wb=openpyxl.load_workbook(SRC)
icons={}   # hash -> datauri
def icon_of(img):
    raw=img._data()
    h=hashlib.md5(raw).hexdigest()[:10]
    if h not in icons:
        im=Image.open(io.BytesIO(raw)).convert('RGBA'); im.thumbnail((40,40))
        b=io.BytesIO(); im.save(b,'WEBP',quality=80)
        icons[h]='data:image/webp;base64,'+base64.b64encode(b.getvalue()).decode()
    return h
def s(v):
    if v is None: return ''
    if isinstance(v,float) and v.is_integer(): v=int(v)
    return str(v).replace('\n',' ').strip()
SCHOOLS=['Aerotheurge','Geomancer','Huntsman','Hydrosophist','Necromancer','Polymorph','Pyrokinetic','Scoundrel','Summoning','Warfare']
skills=[]
for sch in SCHOOLS:
    ws=wb[sch]
    h6={s(ws.cell(6,c).value):c for c in range(1,ws.max_column+1) if ws.cell(6,c).value}
    h7={s(ws.cell(7,c).value):c for c in range(1,ws.max_column+1) if ws.cell(7,c).value}
    col=dict(num=1,name=3,type=4,block=5,bonus=6,eff1=9,eff2=10,dur=h7['Duration'],res=h7['Resistance'],
             self=h7['Target'],ally=h7['Target']+1,enemy=h7['Target']+2,range=h6['Range'],cd=h6['Cool-down'],
             ap=h7['Action'],sp=h7['Source'],mem=h6['Mem. Slots'],book=h6['Skill-book level'],req=h6['Requirements'],sel=h6['Select'])
    imgs={}
    for im in ws._images:
        if im.anchor._from.col==1: imgs[im.anchor._from.row+1]=icon_of(im)
    rows=[r for r in range(9,ws.max_row+1) if isinstance(ws.cell(r,1).value,(int,float))]
    rows.append(ws.max_row+1)
    for i,r0 in enumerate(rows[:-1]):
        r1=rows[i+1]; g=lambda c,r=r0: ws.cell(r,c).value
        sk=dict(school=sch,name=s(g(col['name'])),dmg=s(g(col['type'])),block=s(g(col['block'])),scale=s(g(col['bonus'])),
                range=s(g(col['range'])),cd=s(g(col['cd'])),ap=s(g(col['ap'])),sp=s(g(col['sp'])),mem=s(g(col['mem'])),
                book=s(g(col['book'])),icon=next((imgs[r] for r in range(r0,r1) if r in imgs),None),effects=[],req=[])
        for r in range(r0,r1):
            e=s(ws.cell(r,col['eff2']).value) or s(ws.cell(r,col['eff1']).value)
            if e and e!='-':
                sk['effects'].append(dict(e=e,dur=s(ws.cell(r,col['dur']).value).replace('→','').strip(),
                    chain=('→' in s(ws.cell(r,col['dur']).value)),
                    res=s(ws.cell(r,col['res']).value),
                    t=''.join(k for k,c in (('S',col['self']),('A',col['ally']),('E',col['enemy'])) if s(ws.cell(r,c).value)=='+')))
        if sch=='Warfare':
            mw=s(g(col['req'])); sh=s(g(col['req']+1)); lv=s(g(col['req']+2))
            sk['req'].append(['Warfare',lv])
            if mw=='+': sk['weapon']='Melee weapon'
            if sh=='+': sk['weapon']='Shield'
        elif sch in ('Scoundrel','Huntsman'):
            dg=s(g(col['req'])); lv=s(g(col['req']+1))
            # combo rows may list other abilities
            for r in range(r0,r1):
                a=s(ws.cell(r,col['req']).value); b=s(ws.cell(r,col['req']+1).value)
                if r==r0: sk['req'].append([sch,b]);
                elif a and b: sk['req'].append([a,b])
            if dg=='+': sk['weapon']='Dagger' if sch=='Scoundrel' else 'Bow or crossbow'
        else:
            for r in range(r0,r1):
                a=s(ws.cell(r,col['req']).value); b=s(ws.cell(r,col['req']+1).value)
                if a and b: sk['req'].append([a,b])
        sk['craft']='craft' in sk['book'].lower()
        sk['book']=sk['book'].replace('craft','').strip()
        skills.append(sk)
# statuses
ws=wb['Status Effects']
simg={}
for im in ws._images:
    if im.anchor._from.col==0: simg[im.anchor._from.row+1]=icon_of(im)
starts=[r for r in range(6,ws.max_row+1) if ws.cell(r,2).value]
starts.append(ws.max_row+1)
statuses=[]
for i,r0 in enumerate(starts[:-1]):
    r1=starts[i+1]; st=dict(name=s(ws.cell(r0,2).value),icon=next((simg[r] for r in range(r0,r1) if r in simg),None),effects=[],causes={},resist=[],immune=[],treat=[])
    cat=None; rcat=s(ws.cell(r0,10).value); icat=s(ws.cell(r0,15).value); tcat=s(ws.cell(r0,17).value)
    st['resistHead']=rcat; st['immuneHead']=icat; st['treatHead']=tcat
    for r in range(r0,r1):
        c=s(ws.cell(r,3).value); d=ws.cell(r,4).value
        if c:
            if isinstance(d,(int,float)) and not isinstance(d,bool):
                pct='%' in (ws.cell(r,4).number_format or '')
                dv=(f"{d*100:+g}%" if pct else f"{d:+g}") if d else ''
            else: dv=s(d)
            st['effects'].append((c+(' '+dv if dv else '')).strip())
        e=s(ws.cell(r,5).value)
        if e and r!=r0 or (r==r0 and e): cat=e; st['causes'].setdefault(cat,[])
        f=s(ws.cell(r,6).value)
        if f and cat: st['causes'][cat].append(f)
        k=s(ws.cell(r,11).value); m=s(ws.cell(r,13).value)
        if k: st['resist'].append(k+(' → '+m if m else ''))
        p=s(ws.cell(r,16).value)
        if p: st['immune'].append(p)
        q=s(ws.cell(r,17).value)
        if q and r!=r0: st['treat'].append('§'+q)
        rr=s(ws.cell(r,18).value); u=s(ws.cell(r,21).value)
        if rr: st['treat'].append(rr+(' → '+u if u else ''))
    statuses.append(st)
# special abilities & summons (simple)
def simple(name):
    ws=wb[name]; out=[]; grp=''
    imgs={im.anchor._from.row+1:icon_of(im) for im in ws._images if im.anchor._from.col==1}
    for r in range(9,ws.max_row+1):
        a=ws.cell(r,1).value
        if isinstance(a,str) and a.strip() and not ws.cell(r,3).value: grp=s(a); continue
        b=s(ws.cell(r,2).value); gname=s(ws.cell(r,7).value)
        if b and not isinstance(a,(int,float)): grp=b+(' — '+gname if gname else ''); continue
        if isinstance(a,(int,float)):
            row=[s(ws.cell(r,c).value) for c in range(1,24)]
            out.append(dict(group=grp,name=s(ws.cell(r,3).value),dmg=s(ws.cell(r,4).value),block=s(ws.cell(r,5).value),scale=s(ws.cell(r,6).value),
              effect=s(ws.cell(r,8).value) or s(ws.cell(r,7).value),dur=s(ws.cell(r,10).value),range=s(ws.cell(r,15).value),cd=s(ws.cell(r,16).value),
              ap=s(ws.cell(r,17).value),sp=s(ws.cell(r,18).value),extra=s(ws.cell(r,20).value) if name=='Special Abilities' else '',icon=imgs.get(r)))
    return out
special=simple('Special Abilities'); summons=simple('Summons')
json.dump(dict(skills=skills,statuses=statuses,special=special,summons=summons,icons=icons),open(OUT,'w'))
print(len(skills),len(statuses),len(special),len(summons),len(icons), sum(len(v) for v in icons.values()))
